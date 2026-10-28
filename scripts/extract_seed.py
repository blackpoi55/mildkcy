"""แปลงไฟล์ Excel ต้นฉบับ -> data/seed.json (ใช้ครั้งเดียวตอนนำเข้าข้อมูลเริ่มต้น)

python -I scripts/extract_seed.py "path/to/Motobike ต้นฉบับ.xlsx" data/seed.json
"""
import json
import re
import sys
from datetime import date, timedelta

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

TH_MONTHS = {
    "ม.ค.": 1, "ก.พ.": 2, "มี.ค.": 3, "เม.ย.": 4, "พ.ค.": 5, "มิ.ย.": 6,
    "ก.ค.": 7, "ส.ค.": 8, "ก.ย.": 9, "ต.ค.": 10, "พ.ย.": 11, "ธ.ค.": 12,
}
EN_MONTHS = {m: i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], 1)}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def norm_plate(v):
    return re.sub(r"\s*-\s*", " - ", clean(v))


def th_date(v):
    """'21 มี.ค. 69' / '20 ธ.ค.72' -> '2026-03-21'"""
    s = clean(v)
    if not s:
        return ""
    m = re.match(r"(\d{1,2})\s*([ก-๙\.]+?\.)\s*(\d{2,4})", s)
    if not m:
        return s
    d, mon, y = int(m.group(1)), TH_MONTHS.get(m.group(2).replace(" ", "")), int(m.group(3))
    if not mon:
        return s
    if y < 100:
        y += 2500
    return date(y - 543, mon, d).isoformat()


def norm_time(v):
    s = clean(v)
    if not s:
        return ""
    m = re.match(r"(\d{1,2})[\.:](\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}"
    if re.fullmatch(r"\d{1,2}", s):
        return f"{int(s):02d}:00"
    return s


def main(src, out):
    wb = openpyxl.load_workbook(src)
    cached = openpyxl.load_workbook(src, data_only=True)

    # ---------- ทะเบียนรถ ----------
    ws = wb["Motorbike"]
    bikes = []
    for r in range(2, ws.max_row + 1):
        plate = norm_plate(ws.cell(r, 4).value)
        if not plate:
            continue
        bikes.append({
            "brand": clean(ws.cell(r, 2).value),
            "model": clean(ws.cell(r, 3).value),
            "plate": plate,
            "province": clean(ws.cell(r, 5).value),
            "source": clean(ws.cell(r, 6).value),
            "taxExpiry": th_date(ws.cell(r, 7).value),
            "owner": clean(ws.cell(r, 8).value),
            "cardExpiry": th_date(ws.cell(r, 11).value),
            "transferKit": clean(ws.cell(r, 12).value),
            "transferTo": clean(ws.cell(r, 13).value),
            "status": clean(ws.cell(r, 15).value),
            "dailyRate": 0,
        })
    by_plate = {b["plate"]: b for b in bikes}

    # ---------- ตารางเช่า ----------
    ws = wb["Motorbike Rent"]
    cws = cached["Motorbike Rent"]
    rentals = []
    unknown = {}
    tick_values = {}
    for c0 in range(1, ws.max_column + 1, 24):
        for r0 in range(23, ws.max_row + 1, 22):
            title = clean(ws.cell(r0, c0).value)
            m = re.search(r"([A-Za-z]+)\s+(\d{4})", title)
            if not m:
                continue
            month, year = EN_MONTHS[m.group(1)], int(m.group(2))
            # วันที่ของแต่ละคอลัมน์ อา.-ส.
            days = []
            prev = 0
            cur_m, cur_y = month, year
            for k in range(7):
                v = ws.cell(r0 + 2, c0 + 4 + 2 * k).value
                if isinstance(v, (int, float)):
                    d = int(v)
                    if d < prev:  # ข้ามเดือน
                        cur_m, cur_y = (1, cur_y + 1) if cur_m == 12 else (cur_m + 1, cur_y)
                    prev = d
                    days.append(date(cur_y, cur_m, d))
                else:
                    days.append(None)
            r = r0 + 4
            while r < r0 + 22:
                if clean(ws.cell(r, c0 + 21).value) == "รวมเงิน":
                    break
                plate = norm_plate(ws.cell(r, c0 + 3).value)
                if not plate:
                    r += 1
                    continue
                price = ws.cell(r, c0 + 19).value
                price = float(price) if isinstance(price, (int, float)) else 0
                disc = ws.cell(r, c0 + 20).value
                disc = float(disc) if isinstance(disc, (int, float)) else 0
                fine = ws.cell(r, c0 + 21).value
                fine = float(fine) if isinstance(fine, (int, float)) else 0
                if plate in by_plate:
                    by_plate[plate]["dailyRate"] = price or by_plate[plate]["dailyRate"]
                else:
                    unknown[plate] = (clean(ws.cell(r, c0 + 1).value), clean(ws.cell(r, c0 + 2).value), price)
                first = True
                week_rows = []
                for k, d in enumerate(days):
                    v = ws.cell(r, c0 + 4 + 2 * k).value
                    if d is None or not isinstance(v, (int, float)) or v == 0:
                        continue
                    tick_values[v] = tick_values.get(v, 0) + 1
                    week_rows.append({
                        "date": d.isoformat(),
                        "plate": plate,
                        "time": norm_time(ws.cell(r, c0 + 5 + 2 * k).value),
                        "days": float(v),
                        "price": price,
                        "discount": 0,
                        "fine": 0,
                        "customer": "",
                        "note": "",
                    })
                if not week_rows and (disc or fine):
                    first_day = next(d for d in days if d)
                    week_rows.append({"date": first_day.isoformat(), "plate": plate, "time": "",
                                      "days": 0, "price": price, "discount": 0, "fine": 0,
                                      "customer": "", "note": "ปรับยอด"})
                if week_rows:
                    week_rows[0]["discount"] = disc
                    week_rows[0]["fine"] = fine
                    # ให้ยอดตรงกับที่ Excel คำนวณไว้ (บางแถวสูตรนับวันไม่ตรงกับที่ติ๊ก)
                    excel_w = cws.cell(r, c0 + 22).value
                    mine = sum(x["price"] * x["days"] for x in week_rows) - disc + fine
                    if isinstance(excel_w, (int, float)) and abs(mine - excel_w) > 0.01:
                        week_rows[0]["discount"] += mine - excel_w
                        week_rows[0]["note"] = "ปรับยอดตาม Excel"
                rentals.extend(week_rows)
                r += 1

    # ตรวจยอดรวมกับค่าที่ Excel คำนวณไว้ (แถว 7 / 12)
    total = sum(x["price"] * x["days"] - x["discount"] + x["fine"] for x in rentals)
    excel_total = sum(v for row in (7, 12) for v in
                      (cws.cell(row, c).value for c in range(1, 40))
                      if isinstance(v, (int, float)))
    print(f"bikes={len(bikes)} rentals={len(rentals)} total={total:,.0f} excel={excel_total:,.0f}")
    print("tick values:", tick_values)
    if unknown:
        print("plates not in bike list:", unknown)
        for p, (brand, model, rate) in unknown.items():
            bikes.append({"brand": brand, "model": model, "plate": p, "province": "", "source": "",
                          "taxExpiry": "", "owner": "", "cardExpiry": "", "transferKit": "",
                          "transferTo": "", "status": "เช่า", "dailyRate": rate})

    with open(out, "w", encoding="utf-8") as f:
        json.dump({"bikes": bikes, "rentals": rentals}, f, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
